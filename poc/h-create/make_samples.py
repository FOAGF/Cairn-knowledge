#!/usr/bin/env python
"""Writes the sample files for the M5 prompt check into samples/ (run once; the files are committed).

    python -m venv /tmp/gen && /tmp/gen/bin/pip install reportlab pillow python-docx
    /tmp/gen/bin/python poc/h-create/make_samples.py

All content is made up. MMU and auger are the owner's example terms; Kestrel and P3W are invented
(as in poc/g-terms). Fonts: DejaVu Sans, Noto Sans CJK and URW Z003 (calligraphic, standing in for
handwriting), from the system font folders.
"""
import json, math, os, random

from docx import Document
from docx.shared import Inches
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "samples")
DEJAVU = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
DEJAVU_BOLD = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
CJK = "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"   # index 2 = Simplified Chinese
HAND = "/usr/share/fonts/opentype/urw-base35/Z003-MediumItalic.otf"

# ------------------------------------------------------------------ 1. text PDF: a pump manual
# Filler (cover, contents, legal text, headers, footers) that a draft should drop, and facts it must keep.
MANUAL_PAGES = [
    ("cover", ["Acme Bulk Systems", "Emulsion Pump EP-40", "Service Manual", "Revision C, March 2026",
               "CONFIDENTIAL. For authorised service staff only."]),
    ("contents", ["Contents", "1  Safety notes .......... 3", "2  Calibration .......... 3", "3  Maintenance .......... 4",
                  "", "Legal notice", "Copyright 2026 Acme Bulk Systems. All rights reserved. No part of this manual",
                  "may be reproduced without written permission. Acme accepts no liability for damage caused",
                  "by service work that does not follow this manual."]),
    ("body", ["1  Safety notes",
              "Isolate the pump at the MMU control panel and bleed the line pressure before any service work.",
              "",
              "2  Calibration",
              "Calibrate the flow meter every 250 operating hours, and always after replacing the stator.",
              "1. Flush the line with water until no emulsion comes out.",
              "2. Run water at 120 kg/min for 2 minutes into the calibration tank.",
              "3. Weigh the tank. The weighed amount must be within 2% of the meter reading.",
              "4. If it is not, enter the new K-factor in the MMU controller under Settings > Flow.",
              "   The factory K-factor is 0.982.",
              "WARNING: Never calibrate with emulsion in the line. Emulsion in the calibration tank must be",
              "disposed of as explosive waste."]),
    ("body", ["3  Maintenance",
              "The auger slew bearing and the pump need regular service:",
              "TABLE",
              "",
              "Torque the pump mounting bolts to 210 Nm after every stator change.",
              "Record every service in the MMU logbook."]),
]
MAINT_TABLE = [("Part", "Interval", "Note"),
               ("Auger slew bearing", "weekly", "6 strokes of EP2 grease"),
               ("Pump stator", "1,500 hours", "replace; then calibrate"),
               ("Inlet strainer", "daily", "clean at the end of shift")]


def manual_pdf(path):
    c = canvas.Canvas(path, pagesize=A4)
    c.setTitle("EP-40 Service Manual")
    w, h = A4
    n = len(MANUAL_PAGES)
    for i, (kind, lines) in enumerate(MANUAL_PAGES, 1):
        if i > 1:   # running header and footer, as printed manuals have
            c.setFont("Helvetica", 8)
            c.drawString(50, h - 30, "EP-40 Service Manual, Rev C")
            c.drawRightString(w - 50, 30, f"Page {i} of {n}  |  Confidential  |  Acme Bulk Systems")
        y = h - 90
        if kind == "cover":
            for j, line in enumerate(lines):
                c.setFont("Helvetica-Bold" if j < 3 else "Helvetica", 26 if j == 1 else 14)
                c.drawCentredString(w / 2, h / 2 + 80 - j * 40, line)
        else:
            for line in lines:
                if line == "TABLE":
                    col = [60, 230, 340]
                    for r, row in enumerate(MAINT_TABLE):
                        c.setFont("Helvetica-Bold" if r == 0 else "Helvetica", 10)
                        for x, cell in zip(col, row):
                            c.drawString(x, y, cell)
                        c.line(55, y - 5, 540, y - 5)
                        y -= 20
                    continue
                heading = line[:1].isdigit() and "  " in line[:4] and not line[1:2] == "."
                c.setFont("Helvetica-Bold" if heading or line.startswith(("Contents", "Legal", "WARNING")) else "Helvetica",
                          14 if heading else 10)
                c.drawString(60, y, line)
                y -= 24 if heading else 16
        c.showPage()
    c.save()


# ------------------------------------------------------------------ 2. scanned handover sheet (image)
SCAN_LINES = [
    ("bold", "Shift handover: night shift, 14 Oct"),
    ("text", "Pattern 2214 at P3W: 86 holes loaded, 4 wet holes dewatered first."),
    ("text", "MMU-07: auger hose replaced at 02:30, tested OK."),
    ("text", "Product used: 18.4 t emulsion, 1.2 t AN."),
    ("text", "Misfires: none. Kestrel updated before 6 am."),
    ("cjk", "备注：P3W 北侧台阶有积水，装药前先抽水。"),
    ("hand", "Checked by R. Ortiz, all good"),
]
SCAN_TRUTH = "\n".join(t for _, t in SCAN_LINES)


def scan_image(path):
    random.seed(7)
    W, H = 1400, 1000
    img = Image.new("L", (W, H), 236)
    px = img.load()
    for _ in range(W * H // 25):   # paper grain
        x, y = random.randrange(W), random.randrange(H)
        px[x, y] = random.randint(200, 250)
    d = ImageDraw.Draw(img)
    fonts = {"bold": ImageFont.truetype(DEJAVU_BOLD, 34), "text": ImageFont.truetype(DEJAVU, 28),
             "cjk": ImageFont.truetype(CJK, 30, index=2), "hand": ImageFont.truetype(HAND, 40)}
    y = 70
    for kind, text in SCAN_LINES:
        d.text((70, y), text, font=fonts[kind], fill=30)
        y += 70 if kind == "bold" else 58
    # A sketch of the truck in the margin: a picture the transcription must ignore.
    d.rectangle((1080, 620, 1320, 720), outline=60, width=4)
    d.rectangle((1240, 570, 1320, 620), outline=60, width=4)
    for cx in (1120, 1180, 1290):
        d.ellipse((cx - 22, 705, cx + 22, 749), outline=60, width=4)
    d.line((1080, 640, 1000, 560), fill=60, width=6)
    img = img.rotate(-1.3, resample=Image.BICUBIC, fillcolor=236).filter(ImageFilter.GaussianBlur(0.8))
    img.save(path, quality=80, optimize=True)   # JPEG, like a phone photo or a scanner


# ------------------------------------------------------------------ 3. PDF: one text page + one scanned page
def mixed_pdf(path, scan):
    c = canvas.Canvas(path, pagesize=A4)
    w, h = A4
    c.setFont("Helvetica-Bold", 14)
    c.drawString(60, h - 90, "Site visit notes, 14 Oct")
    c.setFont("Helvetica", 10)
    for i, line in enumerate(["Visited pattern 2214 with the night crew.",
                              "The next page is a scan of the handover sheet they filled in."]):
        c.drawString(60, h - 120 - i * 16, line)
    c.showPage()
    img = ImageReader(scan)   # an image-only page, as a scanner produces
    iw, ih = img.getSize()
    scale = (w - 80) / iw
    c.drawImage(img, 40, h - 60 - ih * scale, iw * scale, ih * scale)
    c.showPage()
    c.save()


# ------------------------------------------------------------------ 4. DOCX: procedure with a table and a picture
def procedure_docx(path, picture):
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Site operations. Controlled document SO-114, version 3"
    doc.add_heading("Pre-start checks for the MMU", level=1)
    doc.add_paragraph("Do these checks at the depot before every shift. Do not leave the depot if any check fails.")
    doc.add_heading("Walk-around", level=2)
    for step in ["Check the AN tank level on the gauge.", "Inspect the auger hoses for cracks and leaks.",
                 "Run the emulsion pump for 30 seconds and read the pressure.",
                 "Check the gassing solution flow at the test port."]:
        doc.add_paragraph(step, style="List Number")
    doc.add_heading("Limits", level=2)
    rows = [("Check", "Limit", "If outside the limit"),
            ("AN tank level", "at least 20%", "refill at the depot"),
            ("Emulsion pump pressure", "8 to 12 bar", "stop and report to the supervisor"),
            ("Auger hoses", "no cracks, no leaks", "replace before leaving"),
            ("Gassing solution flow", "0.8 to 1.2 L/min", "recalibrate the gassing pump")]
    t = doc.add_table(rows=len(rows), cols=3)
    t.style = "Table Grid"
    for r, row in enumerate(rows):
        for col, text in enumerate(row):
            t.cell(r, col).text = text
    doc.add_paragraph("")
    doc.add_picture(picture, width=Inches(2.5))   # a picture the extraction ignores
    doc.add_paragraph("Sign the MMU logbook when all checks pass.")
    doc.save(path)


# ------------------------------------------------------------------ 5. Markdown notes and a task
TWO_TOPICS = """# Notes from the Tuesday meeting

## Flow meter calibration
The EP-40 flow meters on MMU-03 and MMU-07 are due for calibration. Book the calibration tank at the
depot for Thursday. Last time MMU-03 was off by 3.5%, so check it first.

## Travel budget for site visits
Site visits next quarter are capped at four trips. Flights must be booked two weeks ahead, and hotel
receipts go to the admin team within five days of returning.
"""

ONE_TOPIC = """# MMU-07 auger hose leak

The auger on MMU-07 lost pressure on day shift at P3W. The swivel hose at the auger elbow had split
along the outer braid.

## Cause
The hose rubbed on the boom clamp when the auger was fully slewed to the left.

## Fix
- Replaced the swivel hose (part HX-2210).
- Moved the boom clamp 40 mm towards the cab and added a protective sleeve.
- Pressure tested at 12 bar for 5 minutes: no leak.

## Follow-up
Record the repair in Kestrel and check the other trucks' boom clamps at their next PM.
"""

CHINESE = """MMU 冬季保养 记录
气温低于5度时 乳化剂泵启动前要先预热十分钟 否则压力不稳
螺旋输送臂(auger)的软管冬天容易变硬开裂 每周检查一次 发现裂纹立即更换
停车过夜时 把管路里的乳化剂用水冲干净
防冻液每月检查一次 冰点要低于零下20度
"""

TASK = {
    "key": "OPS-41",
    "title": "MMU-07 auger hose keeps cracking",
    "category": "Site operations",
    "description": "The auger swivel hose on MMU-07 has split three times since August. Each time it fails at the "
                   "elbow. Find out why and fix it properly.",
    "comments": [
        ["2026-09-02", "Will look at it tomorrow."],
        ["2026-09-03", "Hose split again at the elbow. Replaced it with the same part (HX-2210). Noticed shiny wear "
                       "marks on the boom clamp."],
        ["2026-09-04", "Thanks!"],
        ["2026-09-10", "Found it: with the auger slewed fully left, the hose rubs on the boom clamp. Moved the clamp "
                       "40 mm towards the cab and added a protective sleeve (part SL-75). Pressure tested at 12 bar for "
                       "5 min, no leak."],
        ["2026-09-24", "Two weeks later, no new wear. MMU-03 and MMU-05 have the same clamp position, so check them "
                       "at their next PM."],
        ["2026-09-25", "ok"],
    ],
}


def main():
    os.makedirs(OUT, exist_ok=True)
    p = lambda name: os.path.join(OUT, name)
    manual_pdf(p("pump-manual.pdf"))
    scan_image(p("handover-scan.jpg"))
    open(p("handover-scan.truth.txt"), "w").write(SCAN_TRUTH + "\n")
    mixed_pdf(p("site-visit-scan.pdf"), p("handover-scan.jpg"))
    procedure_docx(p("pre-start-checks.docx"), p("handover-scan.jpg"))
    open(p("two-topics.md"), "w").write(TWO_TOPICS)
    open(p("one-topic.md"), "w").write(ONE_TOPIC)
    open(p("winter-care-zh.txt"), "w").write(CHINESE)
    json.dump(TASK, open(p("task-ops-41.json"), "w"), ensure_ascii=False, indent=2)
    for f in sorted(os.listdir(OUT)):
        print(f"{os.path.getsize(p(f)):>8}  {f}")


if __name__ == "__main__":
    main()
