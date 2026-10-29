"""M5 additions: file.extract, image paths, finish_reason, streamed usage, timeout and retries,
too_long / bad_file, llm.model_info."""
import base64
import os

import pytest

from cairn_sidecar import ai, errors, extract
from cairn_sidecar.errors import InvalidParams, SidecarError, from_exception
from helpers import FakeReq, Client, target, user
from test_routes import NAMES, make, run

SAMPLES = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "poc", "h-create", "samples")


def sample(name):
    return os.path.join(SAMPLES, name)


def do_extract(path, out_dir, **extra):
    return ai.methods()["file.extract"][0](FakeReq({"path": str(path), "out_dir": str(out_dir), **extra}))


# ---------------------------------------------------------------- file.extract: text
@pytest.mark.parametrize("raw,expected", [
    (b"\xef\xbb\xbfhello \xc3\xa9", "hello é"),                      # UTF-8, BOM dropped
    ("冬季养护".encode("gb18030"), "冬季养护"),  # GB18030
    (b"caf\xe9 \xff", "café ÿ"),                                  # not GB18030 either: Latin-1
])
def test_text_encodings(raw, expected, tmp_path):
    f = tmp_path / "a.txt"
    f.write_bytes(raw)
    res = do_extract(f, tmp_path / "out")
    assert res == {"kind": "text", "pages": 1, "parts": [{"page": 1, "text": expected}]}
    assert not (tmp_path / "out").exists(), "nothing to write for text"


def test_markdown_sample(tmp_path):
    res = do_extract(sample("one-topic.md"), tmp_path / "o")
    assert res["kind"] == "text" and res["parts"][0]["text"].startswith("#")
    res = do_extract(sample("winter-care-zh.txt"), tmp_path / "o")
    assert any("一" <= ch <= "鿿" for ch in res["parts"][0]["text"])


# ---------------------------------------------------------------- file.extract: PDF
def test_pdf_text_pages(tmp_path):
    res = do_extract(sample("pump-manual.pdf"), tmp_path / "o")
    assert res["kind"] == "pdf" and res["pages"] == 4
    assert [p["page"] for p in res["parts"]] == [1, 2, 3, 4]
    assert all("text" in p and "image_path" not in p for p in res["parts"])
    assert "\r" not in "".join(p["text"] for p in res["parts"])


def test_pdf_scanned_page_becomes_a_jpeg(tmp_path):
    from PIL import Image
    out = tmp_path / "o"
    res = do_extract(sample("site-visit-scan.pdf"), out)
    assert res["pages"] == 2
    assert "Site visit notes" in res["parts"][0]["text"]
    part = res["parts"][1]
    assert part["page"] == 2 and part["image_path"] == str(out / "p0002.jpg")
    with Image.open(part["image_path"]) as im:
        assert im.format == "JPEG" and im.mode == "L"
        assert im.size == (part["width"], part["height"]) and max(im.size) == 2000


def test_pdf_dpi_and_max_edge(tmp_path):
    res = do_extract(sample("site-visit-scan.pdf"), tmp_path / "o", max_edge=500, dpi=72)
    assert max(res["parts"][1]["width"], res["parts"][1]["height"]) <= 500
    big = do_extract(sample("site-visit-scan.pdf"), tmp_path / "o2", max_edge=5000, dpi=72)
    assert max(big["parts"][1]["width"], big["parts"][1]["height"]) < 1000  # 72 dpi, A4: about 595 x 842


def test_pdf_damaged_or_not_a_pdf_is_bad_file(tmp_path):
    f = tmp_path / "x.pdf"
    f.write_bytes(b"this is not a pdf")
    with pytest.raises(SidecarError) as ei:
        do_extract(f, tmp_path / "o")
    assert ei.value.kind == "bad_file" and ei.value.data()["kind"] == "bad_file"
    g = tmp_path / "y.pdf"
    g.write_bytes(open(sample("pump-manual.pdf"), "rb").read()[:300])
    with pytest.raises(SidecarError) as ei:
        do_extract(g, tmp_path / "o")
    assert ei.value.kind == "bad_file"


def test_pdf_password_error_is_bad_file(tmp_path, monkeypatch):
    """pypdfium2 raises PdfiumError (\"Incorrect password\") for an encrypted file."""
    import pypdfium2 as pdfium

    def locked(path):
        raise pdfium.PdfiumError("Failed to load document (PDFium: Incorrect password error).")
    monkeypatch.setattr(pdfium, "PdfDocument", locked)
    with pytest.raises(SidecarError) as ei:
        do_extract(sample("pump-manual.pdf"), tmp_path / "o")
    assert ei.value.kind == "bad_file" and "password" in ei.value.message


# ---------------------------------------------------------------- file.extract: DOCX
def test_docx_sample_markdown(tmp_path):
    res = do_extract(sample("pre-start-checks.docx"), tmp_path / "o")
    assert res["kind"] == "docx" and res["pages"] == 1
    text = res["parts"][0]["text"]
    lines = text.split("\n")
    assert lines[0] == "# Pre-start checks for the MMU"
    assert "## Walk-around" in lines and "## Limits" in lines
    assert "1. Check the AN tank level on the gauge." in lines and "4. Check the gassing solution flow at the test port." in lines
    assert "| Check | Limit | If outside the limit |" in lines and "|---|---|---|" in lines
    assert lines[-1] == "Sign the MMU logbook when all checks pass."


def test_docx_styles_escapes_and_skips(tmp_path):
    import docx
    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "HEADER TEXT"
    d.sections[0].footer.paragraphs[0].text = "FOOTER TEXT"
    d.add_heading("Top", 1)
    d.add_paragraph("one", style="List Bullet")
    d.add_paragraph("two", style="List Bullet")
    d.add_paragraph("first", style="List Number")
    d.add_paragraph("second", style="List Number")
    d.add_heading("Sub", 3)
    d.add_paragraph("plain text")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "a|b", "c"
    t.cell(1, 0).text, t.cell(1, 1).text = "line1\nline2", ""
    f = tmp_path / "s.docx"
    d.save(f)
    text = do_extract(f, tmp_path / "o")["parts"][0]["text"]
    assert text == ("# Top\n\n- one\n- two\n\n1. first\n2. second\n\n### Sub\n\nplain text\n\n"
                    "| a\\|b | c |\n|---|---|\n| line1 line2 |  |")
    assert "HEADER" not in text and "FOOTER" not in text


def test_docx_damaged_is_bad_file(tmp_path):
    f = tmp_path / "x.docx"
    f.write_bytes(b"PK not really a zip")
    with pytest.raises(SidecarError) as ei:
        do_extract(f, tmp_path / "o")
    assert ei.value.kind == "bad_file"


# ---------------------------------------------------------------- file.extract: images
def test_image_sample_is_scaled_jpeg(tmp_path):
    from PIL import Image
    out = tmp_path / "o"
    res = do_extract(sample("handover-scan.jpg"), out)
    assert res["kind"] == "image" and res["pages"] == 1
    part = res["parts"][0]
    assert part["page"] == 1 and part["image_path"] == str(out / "image.jpg")
    with Image.open(part["image_path"]) as im:
        assert im.format == "JPEG" and im.size == (part["width"], part["height"]) and max(im.size) <= 2000


def test_image_exif_upright_rgb_and_scaling(tmp_path):
    from PIL import Image
    src = Image.new("RGBA", (3000, 1000), (255, 0, 0, 128))
    f = tmp_path / "wide.png"
    src.save(f)
    res = do_extract(f, tmp_path / "o")
    p = res["parts"][0]
    assert (p["width"], p["height"]) == (2000, 667)
    with Image.open(p["image_path"]) as im:
        assert im.mode == "RGB"
    # EXIF orientation 6: stored 300x100, shown 100x300
    exif = Image.Exif()
    exif[0x0112] = 6
    g = tmp_path / "rot.jpg"
    Image.new("RGB", (300, 100), (0, 255, 0)).save(g, exif=exif)
    p = do_extract(g, tmp_path / "o2")["parts"][0]
    assert (p["width"], p["height"]) == (100, 300)
    gray = tmp_path / "gray.png"
    Image.new("L", (40, 20), 100).save(gray)
    with Image.open(do_extract(gray, tmp_path / "o3")["parts"][0]["image_path"]) as im:
        assert im.mode == "L"
    p = do_extract(f, tmp_path / "o4", max_edge=1000)["parts"][0]
    assert (p["width"], p["height"]) == (1000, 333)


def test_image_damaged_is_bad_file(tmp_path):
    f = tmp_path / "x.png"
    f.write_bytes(b"not an image")
    with pytest.raises(SidecarError) as ei:
        do_extract(f, tmp_path / "o")
    assert ei.value.kind == "bad_file"


def test_extract_param_errors(tmp_path):
    with pytest.raises(InvalidParams):
        do_extract(tmp_path / "a.xyz", tmp_path / "o")
    with pytest.raises(InvalidParams):
        ai.methods()["file.extract"][0](FakeReq({"path": "a.txt"}))
    f = tmp_path / "a.txt"
    f.write_text("x")
    with pytest.raises(InvalidParams):
        do_extract(f, tmp_path / "o", max_edge=0)
    with pytest.raises(SidecarError) as ei:
        do_extract(tmp_path / "missing.pdf", tmp_path / "o")
    assert ei.value.kind == "bad_file"


def test_extract_over_stdio_and_libraries_stay_lazy(tmp_path):
    c = Client("ai")
    try:
        before = c.call("ping")["result"]
        assert before["pong"]
        res = c.call("file.extract", {"path": sample("site-visit-scan.pdf"), "out_dir": str(tmp_path / "o")})
        assert res["result"]["parts"][1]["image_path"].endswith("p0002.jpg")
        err = c.call("file.extract", {"path": str(tmp_path / "nope.pdf"), "out_dir": str(tmp_path / "o")})["error"]
        assert err["data"]["kind"] == "bad_file"
    finally:
        c.close()


def test_extract_libraries_not_imported_at_start():
    import subprocess
    import sys
    from conftest import clean_env
    p = subprocess.Popen([sys.executable, "-I", "-X", "importtime", "-m", "cairn_sidecar", "--role", "ai"],
                         stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True,
                         env=clean_env())
    out, err = p.communicate('{"jsonrpc":"2.0","id":1,"method":"ping","params":{}}\n', timeout=20)
    imported = {line.rsplit("|", 1)[-1].strip().split(".")[0] for line in err.splitlines() if "|" in line}
    for heavy in ("pypdfium2", "docx", "PIL", "lxml"):
        assert heavy not in imported, f"{heavy} was imported at start"


# ---------------------------------------------------------------- image blocks by path
@pytest.mark.parametrize("stream", [False, True], ids=["nostream", "stream"])
@pytest.mark.parametrize("name", NAMES)
def test_image_path_block_is_read_into_a_data_url(name, stream, mock, tmp_path):
    f = tmp_path / "page.jpg"
    f.write_bytes(b"\xff\xd8\xff fake jpeg bytes")
    b64 = base64.b64encode(f.read_bytes()).decode()
    msgs = [{"role": "user", "content": [{"type": "text", "text": "read"}, {"type": "image", "path": str(f)}]}]
    t, *_ = make(mock, name)
    run(t, messages=msgs, stream=stream, on_delta=lambda d: None)
    body = mock.requests[-1]["body"]
    assert b64 in str(body) and str(f) not in str(body)
    if t["format"] == "anthropic_messages":
        img = [b for b in body["messages"][0]["content"] if b["type"] == "image"][0]
        assert img["source"]["media_type"] == "image/jpeg" and img["source"]["data"] == b64
    elif t["format"] == "chat_completions":
        urls = [b["image_url"]["url"] for b in body["messages"][0]["content"] if b["type"] == "image_url"]
        assert urls == [f"data:image/jpeg;base64,{b64}"]
    else:
        urls = [b["image_url"] for b in body["input"][0]["content"] if b["type"] == "input_image"]
        assert urls == [f"data:image/jpeg;base64,{b64}"]


def test_image_path_mime_by_extension_and_errors(tmp_path):
    for ext, mime in ((".png", "image/png"), (".JPEG", "image/jpeg"), (".webp", "image/webp"), (".gif", "image/gif")):
        f = tmp_path / f"a{ext}"
        f.write_bytes(b"x")
        assert ai._image_url({"type": "image", "path": str(f)}).startswith(f"data:{mime};base64,")
    assert ai._image_url({"type": "image", "data_url": "data:image/png;base64,AA=="}) == "data:image/png;base64,AA=="
    with pytest.raises(SidecarError) as ei:
        ai._image_url({"type": "image", "path": str(tmp_path / "missing.png")})
    assert ei.value.kind == "bad_file"
    with pytest.raises(InvalidParams):
        ai._image_url({"type": "image", "path": str(tmp_path / "a.bmp")})
    with pytest.raises(InvalidParams):
        ai._image_url({"type": "image"})


# ---------------------------------------------------------------- finish_reason
@pytest.mark.parametrize("stream", [False, True], ids=["nostream", "stream"])
@pytest.mark.parametrize("name", NAMES)
def test_finish_reason_stop_and_length(name, stream, mock):
    t, *_ = make(mock, name)
    assert run(t, stream=stream, on_delta=lambda d: None)["finish_reason"] == "stop"
    mock.reset("length")
    r = run(t, stream=stream, on_delta=lambda d: None)
    assert r["finish_reason"] == "length" and r["text"] == "Hello from mock"


def test_finish_reason_mapping():
    assert [ai._finish(x) for x in ("stop", "end_turn", "stop_sequence", "length", "max_tokens", "max_output_tokens")] \
        == ["stop", "stop", "stop", "length", "length", "length"]
    assert ai._finish(None) is None and ai._finish("content_filter") is None


def test_llm_complete_reports_finish_reason(mock):
    mock.reset("length")
    r = FakeReq({"target": target(endpoint=mock.url + "/v1"), "messages": user(), "max_tokens": 5})
    assert ai.llm_complete(r)["finish_reason"] == "length"


def test_cancelled_stream_has_finish_reason_cancelled(mock):
    import threading
    t = ai.Target(target(endpoint=mock.url + "/v1"))
    ev = threading.Event()
    ev.set()
    r = ai.run_completion(t, user(), stream=True, cancelled=ev)
    assert r["finish_reason"] == "cancelled"
    req = FakeReq({"target": target(endpoint=mock.url + "/v1"), "messages": user()})
    ai.llm_complete(req)
    assert req.cancel_result()["finish_reason"] == "cancelled"


def test_cancel_over_stdio_says_cancelled(new_mock):
    import time
    slow = new_mock()
    slow.reset("slow")
    c = Client("ai")
    try:
        rid = c.send("llm.complete", {"target": target(endpoint=slow.url + "/v1"), "messages": user(), "stream": True})
        assert c.next()["method"] == "stream"
        c.send("cancel", {"id": rid}, notify=True)
        res = c.wait_for(rid, timeout=5)[1]["result"]
        assert res["finish_reason"] == "cancelled" and res["cancelled"] is True
        time.sleep(0.1)
    finally:
        c.close()


# ---------------------------------------------------------------- usage on streamed calls
@pytest.mark.parametrize("name", NAMES)
def test_streamed_calls_report_usage_on_every_route(name, mock):
    t, *_ = make(mock, name)
    r = run(t, stream=True, on_delta=lambda d: None)
    assert r["usage"] == {"input_tokens": 5, "output_tokens": 3}
    body = mock.requests[-1]["body"]
    if t["format"] == "chat_completions":
        assert body["stream_options"] == {"include_usage": True}


# ---------------------------------------------------------------- timeout and retries
def test_timeout_s_is_passed_per_request(mock):
    t = ai.Target(target(endpoint=mock.url + "/v1"))
    assert ai.build_call(t, user())[1]["timeout"] == ai.TIMEOUT == 90
    api, kw = ai.build_call(t, user(), timeout=12)
    assert kw["timeout"] == 12 and kw["num_retries"] == 0, "LiteLLM does not retry; the sidecar does"
    for bad in (0, -1, "9", True):
        with pytest.raises(InvalidParams):
            ai.llm_complete(FakeReq({"target": target(endpoint=mock.url + "/v1"), "messages": user(), "timeout_s": bad}))


def test_timeout_is_retried_once_then_succeeds(mock, monkeypatch):
    monkeypatch.setattr(ai, "NUM_RETRIES", 3)
    mock.reset("hang_then_ok")
    r = FakeReq({"target": target(endpoint=mock.url + "/v1"), "messages": user(), "timeout_s": 0.4})
    assert ai.llm_complete(r)["text"] == "Hello from mock"
    assert len(mock.requests) == 2


def test_timeout_fails_after_one_retry_even_with_three_retries_allowed(mock, monkeypatch):
    monkeypatch.setattr(ai, "NUM_RETRIES", 3)
    mock.reset("hang")
    r = FakeReq({"target": target(endpoint=mock.url + "/v1"), "messages": user(), "timeout_s": 0.4})
    with pytest.raises(Exception) as ei:
        ai.llm_complete(r)
    err = from_exception(ei.value, ["SECRET"])
    assert err.kind == "network"
    assert len(mock.requests) == 2, "a timeout is retried at most once"


def test_rate_limit_retried_three_times(mock, monkeypatch):
    monkeypatch.setattr(ai, "NUM_RETRIES", 3)
    mock.reset("429")
    t, *_ = make(mock, "openai_compat_chat")
    with pytest.raises(Exception) as ei:
        run(t)
    assert from_exception(ei.value).kind == "rate_limit"
    assert len(mock.requests) == 4


class APIConnectionError(Exception):
    pass


class Timeout(Exception):
    pass


class RateLimitError(Exception):
    status_code = 429


class AuthenticationError(Exception):
    status_code = 401


def attempts(excs, retries=3):
    """_with_retries over a function that raises `excs` in turn, then returns "ok"."""
    calls = []

    def attempt():
        calls.append(1)
        if len(calls) <= len(excs):
            raise excs[len(calls) - 1]
        return "ok"
    return attempt, calls


def test_retry_rule_in_isolation(monkeypatch):
    monkeypatch.setattr(ai, "NUM_RETRIES", 3)
    sent = lambda: False  # noqa: E731
    f, calls = attempts([APIConnectionError("x"), RateLimitError("y"), APIConnectionError("z")])
    assert ai._with_retries(f, sent) == "ok" and len(calls) == 4
    f, calls = attempts([APIConnectionError("x")] * 4)
    with pytest.raises(APIConnectionError):
        ai._with_retries(f, sent)
    assert len(calls) == 4  # first call + 3 retries
    f, calls = attempts([Timeout("t"), Timeout("t")])
    with pytest.raises(Timeout):
        ai._with_retries(f, sent)
    assert len(calls) == 2  # first call + 1 retry
    f, calls = attempts([Timeout("t")])
    assert ai._with_retries(f, sent) == "ok" and len(calls) == 2
    # a timeout and then connection errors: the timeout budget is separate
    f, calls = attempts([Timeout("t"), APIConnectionError("c"), APIConnectionError("c")])
    assert ai._with_retries(f, sent) == "ok" and len(calls) == 4
    # other errors are never retried
    f, calls = attempts([AuthenticationError("no")])
    with pytest.raises(AuthenticationError):
        ai._with_retries(f, sent)
    assert len(calls) == 1


def test_no_retry_after_stream_text_was_sent(monkeypatch):
    monkeypatch.setattr(ai, "NUM_RETRIES", 3)
    f, calls = attempts([APIConnectionError("dropped")])
    with pytest.raises(APIConnectionError):
        ai._with_retries(f, lambda: True)
    assert len(calls) == 1
    # but a stream that failed before any text is retried
    f, calls = attempts([APIConnectionError("dropped")])
    assert ai._with_retries(f, lambda: False) == "ok" and len(calls) == 2


def test_no_retry_when_cancelled():
    import threading
    ev = threading.Event()
    ev.set()
    f, calls = attempts([APIConnectionError("x")])
    with pytest.raises(APIConnectionError):
        ai._with_retries(f, lambda: False, ev)
    assert len(calls) == 1


def test_backoff_doubles(monkeypatch):
    monkeypatch.setattr(ai, "NUM_RETRIES", 3)
    monkeypatch.setattr(ai, "BACKOFF", 0.5)
    waits = []
    monkeypatch.setattr(ai, "_sleep", lambda s, c: waits.append(s))
    f, calls = attempts([APIConnectionError("x")] * 3)
    assert ai._with_retries(f, lambda: False) == "ok"
    assert waits == [0.5, 1.0, 2.0]


def test_is_timeout():
    assert errors.is_timeout(Timeout("x")) and not errors.is_timeout(APIConnectionError("x"))
    e = APIConnectionError("Connection error: timed out")
    assert errors.is_timeout(e)


# ---------------------------------------------------------------- too_long
@pytest.mark.parametrize("name", NAMES)
def test_too_long_from_a_400(name, mock):
    mock.reset("too_long")
    t, *_ = make(mock, name)
    with pytest.raises(Exception) as ei:
        run(t)
    err = from_exception(ei.value, ["SECRET"])
    assert err.kind == "too_long" and err.status == 400
    assert err.data()["hint"]
    assert len(mock.requests) == 1


def test_too_long_classification():
    class ContextWindowExceededError(Exception):
        status_code = 400

    class BadRequestError(Exception):
        status_code = 400
    assert from_exception(ContextWindowExceededError("whatever")).kind == "too_long"
    for msg in ("prompt is too long: 300000 tokens > 200000 maximum",
                "The input is too long for this model",
                "This model's maximum context length is 8192 tokens",
                "Input exceeds the context window of this model",
                "request exceeds the maximum context length",
                "Error code: context_length_exceeded"):
        assert from_exception(BadRequestError(msg)).kind == "too_long", msg
    assert from_exception(BadRequestError("reasoning_effort is not supported")).kind == "bad_request"
    assert from_exception(BadRequestError("temperature too high")).kind == "bad_request"
    # only a 400 counts: a 500 that mentions it stays internal
    e = Exception("the prompt is too long")
    e.status_code = 500
    assert from_exception(e).kind == "internal"


def test_bad_file_is_an_error_kind():
    err = extract.bad_file("x")
    assert err.data()["kind"] == "bad_file" and err.data()["hint"]
    assert errors.MESSAGES["bad_file"]


# ---------------------------------------------------------------- llm.model_info
def info(**over):
    return ai.llm_model_info(FakeReq({"target": target(**over)}))


def test_model_info_known_model_from_local_map():
    res = info(provider_kind="anthropic", format="anthropic_messages", model="claude-sonnet-4-5")
    assert set(res) == {"max_input_tokens", "max_output_tokens"}
    assert isinstance(res["max_input_tokens"], int) and res["max_input_tokens"] >= 100000
    assert isinstance(res["max_output_tokens"], int) and res["max_output_tokens"] > 1000


def test_model_info_unknown_model_is_null():
    assert info(model="my-deployment-xyz") == {"max_input_tokens": None, "max_output_tokens": None}


def test_model_info_uses_the_route_not_just_the_name(monkeypatch):
    seen = []

    def fake(model, **kw):
        seen.append(model)
        return {"max_input_tokens": 1000, "max_output_tokens": None}
    monkeypatch.setattr(ai.litellm(), "get_model_info", fake)
    assert info(provider_kind="azure_foundry", endpoint="https://x/openai", model="gpt-5") == \
        {"max_input_tokens": 1000, "max_output_tokens": None}
    info(provider_kind="azure_foundry", endpoint="https://x/openai/v1", model="gpt-5")
    info(provider_kind="azure_foundry", endpoint="https://x/anthropic", format="anthropic_messages", model="claude-x")
    info(provider_kind="azure_foundry", endpoint="https://x/anthropic", format="anthropic_messages", model="my-sonnet")
    info(provider_kind="anthropic", format="anthropic_messages", model="claude-x")
    info(provider_kind="local", model="llama")
    assert seen == ["azure/gpt-5", "openai/gpt-5", "azure_ai/claude-x", "anthropic/my-sonnet", "anthropic/claude-x",
                    "openai/llama"]


def test_model_info_exception_and_odd_values_are_null(monkeypatch):
    def boom(model, **kw):
        raise RuntimeError("no")
    monkeypatch.setattr(ai.litellm(), "get_model_info", boom)
    assert info() == {"max_input_tokens": None, "max_output_tokens": None}
    monkeypatch.setattr(ai.litellm(), "get_model_info", lambda model, **kw: {"max_input_tokens": 0, "max_output_tokens": "x"})
    assert info() == {"max_input_tokens": None, "max_output_tokens": None}


def test_model_info_makes_no_network_call(mock):
    info(endpoint=mock.url + "/v1")
    assert mock.requests == []


def test_model_info_over_stdio():
    c = Client("ai")
    try:
        res = c.call("llm.model_info", {"target": target(model="nobody-knows-me")})
        assert res["result"] == {"max_input_tokens": None, "max_output_tokens": None}
    finally:
        c.close()
